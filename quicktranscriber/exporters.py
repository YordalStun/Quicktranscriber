"""Export meetings as TXT, Markdown, Word, subtitles, JSON, HTML or a ZIP of everything."""

from __future__ import annotations

import html
import io
import json
import re
import zipfile
from datetime import datetime
from typing import Any

from . import __version__, db, paths
from .pipeline import speaker_names

FORMATS = {
    "txt": ("Plain text", "text/plain; charset=utf-8"),
    "md": ("Markdown", "text/markdown; charset=utf-8"),
    "docx": ("Word document", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "html": ("Web page / PDF", "text/html; charset=utf-8"),
    "srt": ("Subtitles (SRT)", "application/x-subrip; charset=utf-8"),
    "vtt": ("Subtitles (WebVTT)", "text/vtt; charset=utf-8"),
    "json": ("JSON data", "application/json; charset=utf-8"),
    "zip": ("Everything (ZIP)", "application/zip"),
}


def ts(seconds: float, always_hours: bool = True) -> str:
    s = max(0, int(seconds))
    h, m, sec = s // 3600, s % 3600 // 60, s % 60
    return f"{h:02d}:{m:02d}:{sec:02d}" if always_hours or h else f"{m:02d}:{sec:02d}"


def gather(meeting_id: str) -> dict[str, Any]:
    m = db.one("SELECT * FROM meetings WHERE id = ?", (meeting_id,))
    if not m:
        raise KeyError(meeting_id)
    names = speaker_names(meeting_id)
    speakers = db.query(
        "SELECT ms.key, ms.color, ms.talk_time, s.name FROM meeting_speakers ms "
        "LEFT JOIN speakers s ON s.id = ms.speaker_id WHERE ms.meeting_id = ? ORDER BY ms.talk_time DESC",
        (meeting_id,),
    )
    segments = db.query(
        "SELECT start, end, speaker, text, words FROM segments WHERE meeting_id = ? ORDER BY idx", (meeting_id,)
    )
    for s in segments:
        s["name"] = names.get(s["speaker"], s["speaker"] or "")
        s["words"] = db.loads(s["words"], [])
    used = {s["speaker"] for s in segments}
    return {
        "id": m["id"],
        "title": m["title"],
        "date": datetime.fromtimestamp(m["recorded_at"] or m["created_at"]),
        "duration": m["duration"] or 0,
        "language": m["language"],
        "notes": db.loads(m["notes"], None),
        "notes_meta": db.loads(m["notes_meta"], {}),
        "speakers": [
            {"key": s["key"], "name": names.get(s["key"], s["key"]), "color": s["color"], "talk_time": s["talk_time"]}
            for s in speakers if s["key"] in used
        ],
        "segments": segments,
    }


def filename(data: dict[str, Any], ext: str) -> str:
    safe = re.sub(r"[^\w\- ]+", "", data["title"]).strip().replace(" ", "_")[:60] or "meeting"
    return f"{data['date']:%Y-%m-%d}_{safe}.{ext}"


def _meta_line(d: dict[str, Any]) -> str:
    who = ", ".join(s["name"] for s in d["speakers"])
    parts = [f"{d['date']:%A %d %B %Y, %H:%M}", f"Duration {ts(d['duration'])}"]
    if who:
        parts.append(f"Speakers: {who}")
    return " · ".join(parts)


# ---------------------------------------------------------------------------
# Text formats
# ---------------------------------------------------------------------------


def _notes_blocks(n: dict[str, Any]) -> list[tuple[str, Any]]:
    blocks: list[tuple[str, Any]] = []
    if n.get("summary"):
        blocks.append(("Summary", n["summary"]))
    if n.get("key_points"):
        blocks.append(("Key points", n["key_points"]))
    if n.get("decisions"):
        blocks.append(("Decisions", n["decisions"]))
    if n.get("action_items"):
        blocks.append(("Action items", n["action_items"]))
    if n.get("open_questions"):
        blocks.append(("Open questions", n["open_questions"]))
    for s in n.get("sections") or []:
        blocks.append((s["title"], s["items"]))
    if n.get("chapters"):
        blocks.append(("Chapters", n["chapters"]))
    return blocks


def _action_text(a: dict[str, Any], timestamps: bool) -> str:
    extra = []
    if a.get("owner"):
        extra.append(a["owner"])
    if a.get("due"):
        extra.append(f"due {a['due']}")
    text = a["task"] + (f" ({'; '.join(extra)})" if extra else "")
    if timestamps and a.get("t") is not None:
        text += f" [{ts(a['t'])}]"
    return text


def to_txt(d: dict[str, Any], notes: bool = True, transcript: bool = True, timestamps: bool = True) -> str:
    out = [d["title"].upper(), _meta_line(d), "=" * 72, ""]
    n = d["notes"]
    if notes and n:
        for title, body in _notes_blocks(n):
            out.append(title.upper())
            if isinstance(body, str):
                out.append(body)
            elif title == "Action items":
                out += [("[x] " if a.get("done") else "[ ] ") + _action_text(a, timestamps) for a in body]
            elif title == "Chapters":
                out += [(f"[{ts(c['t'])}] " if timestamps else "") + c["title"] + (f" - {c['summary']}" if c.get("summary") else "") for c in body]
            else:
                out += [f"  • {x}" for x in body]
            out.append("")
    if transcript and d["segments"]:
        out += ["TRANSCRIPT", "-" * 72]
        for s in d["segments"]:
            prefix = f"[{ts(s['start'])}] " if timestamps else ""
            out.append(f"{prefix}{s['name']}: {s['text']}")
            out.append("")
    out.append(f"Created with QuickTranscriber {__version__} - processed privately on this computer.")
    return "\n".join(out).strip() + "\n"


def to_md(d: dict[str, Any], notes: bool = True, transcript: bool = True, timestamps: bool = True) -> str:
    out = [f"# {d['title']}", "", f"*{_meta_line(d)}*", ""]
    n = d["notes"]
    if notes and n:
        for title, body in _notes_blocks(n):
            out.append(f"## {title}")
            out.append("")
            if isinstance(body, str):
                out.append(body)
            elif title == "Action items":
                out += [("- [x] " if a.get("done") else "- [ ] ") + _action_text(a, timestamps) for a in body]
            elif title == "Chapters":
                out += [f"- " + (f"**{ts(c['t'])}** " if timestamps else "") + c["title"] + (f" — {c['summary']}" if c.get("summary") else "") for c in body]
            else:
                out += [f"- {x}" for x in body]
            out.append("")
    if transcript and d["segments"]:
        out += ["## Transcript", ""]
        for s in d["segments"]:
            prefix = f"`{ts(s['start'])}` " if timestamps else ""
            out.append(f"{prefix}**{s['name']}:** {s['text']}")
            out.append("")
    return "\n".join(out).strip() + "\n"


def _cues(d: dict[str, Any], max_seconds: float = 6.0, max_chars: int = 84) -> list[tuple[float, float, str, str]]:
    cues = []
    for s in d["segments"]:
        words = s["words"] or [[s["start"], s["end"], " " + s["text"], 1]]
        cur: list = []
        for w in words:
            text_len = len("".join(x[2] for x in cur).strip())
            if cur and (w[1] - cur[0][0] > max_seconds or text_len + len(w[2]) > max_chars):
                cues.append((cur[0][0], cur[-1][1], "".join(x[2] for x in cur).strip(), s["name"]))
                cur = []
            cur.append(w)
        if cur:
            cues.append((cur[0][0], cur[-1][1], "".join(x[2] for x in cur).strip(), s["name"]))
    return cues


def _srt_time(t: float) -> str:
    ms = int(round(t * 1000))
    return f"{ms // 3600000:02d}:{ms % 3600000 // 60000:02d}:{ms % 60000 // 1000:02d},{ms % 1000:03d}"


def to_srt(d: dict[str, Any], speakers: bool = True, **_: Any) -> str:
    out = []
    for i, (a, b, text, who) in enumerate(_cues(d), 1):
        out += [str(i), f"{_srt_time(a)} --> {_srt_time(max(b, a + 0.3))}", f"{who}: {text}" if speakers else text, ""]
    return "\n".join(out)


def to_vtt(d: dict[str, Any], speakers: bool = True, **_: Any) -> str:
    out = ["WEBVTT", ""]
    for a, b, text, who in _cues(d):
        out += [f"{_srt_time(a).replace(',', '.')} --> {_srt_time(max(b, a + 0.3)).replace(',', '.')}",
                f"<v {who}>{text}" if speakers else text, ""]
    return "\n".join(out)


def to_json(d: dict[str, Any], **_: Any) -> str:
    data = dict(d)
    data["date"] = d["date"].isoformat()
    data["app"] = f"QuickTranscriber {__version__}"
    return json.dumps(data, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------


def to_docx(d: dict[str, Any], notes: bool = True, transcript: bool = True, timestamps: bool = True) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: F401
    from docx.shared import Pt, RGBColor

    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    doc.add_heading(d["title"], level=0)
    meta = doc.add_paragraph(_meta_line(d))
    meta.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    n = d["notes"]
    if notes and n:
        for title, body in _notes_blocks(n):
            doc.add_heading(title, level=1)
            if isinstance(body, str):
                for para in body.split("\n"):
                    if para.strip():
                        doc.add_paragraph(para.strip())
            elif title == "Action items":
                table = doc.add_table(rows=1, cols=4 if timestamps else 3)
                table.style = "Light Grid Accent 1"
                hdr = table.rows[0].cells
                for i, h in enumerate(["Task", "Owner", "Due"] + (["Time"] if timestamps else [])):
                    hdr[i].text = h
                for a in body:
                    row = table.add_row().cells
                    row[0].text = ("✓ " if a.get("done") else "") + a["task"]
                    row[1].text = a.get("owner") or ""
                    row[2].text = a.get("due") or ""
                    if timestamps:
                        row[3].text = ts(a["t"]) if a.get("t") is not None else ""
            elif title == "Chapters":
                for c in body:
                    p = doc.add_paragraph(style="List Bullet")
                    if timestamps:
                        r = p.add_run(ts(c["t"]) + "  ")
                        r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
                    p.add_run(c["title"]).bold = True
                    if c.get("summary"):
                        p.add_run(" — " + c["summary"])
            else:
                for x in body:
                    doc.add_paragraph(x, style="List Bullet")
    if transcript and d["segments"]:
        doc.add_page_break()
        doc.add_heading("Transcript", level=1)
        colors = {s["key"]: s["color"] for s in d["speakers"]}
        for s in d["segments"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            if timestamps:
                r = p.add_run(ts(s["start"]) + "  ")
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
            r = p.add_run(s["name"] + ": ")
            r.bold = True
            col = colors.get(s["speaker"])
            if col and re.fullmatch(r"#[0-9a-fA-F]{6}", col):
                r.font.color.rgb = RGBColor.from_string(col[1:].upper())
            p.add_run(s["text"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# HTML (also for printing to PDF)
# ---------------------------------------------------------------------------

HTML_CSS = """
:root{--ink:#111827;--muted:#6b7280;--line:#e5e7eb;--accent:#7c3aed}
*{box-sizing:border-box}body{font:15px/1.6 -apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif;color:var(--ink);
max-width:860px;margin:40px auto;padding:0 24px;background:#fff}
h1{font-size:30px;line-height:1.2;margin:0 0 6px}h2{font-size:18px;margin:32px 0 10px;padding-bottom:6px;
border-bottom:2px solid var(--line)}.meta{color:var(--muted);margin-bottom:24px}
ul{padding-left:22px}li{margin:4px 0}.t{color:var(--muted);font-variant-numeric:tabular-nums;font-size:13px}
table{border-collapse:collapse;width:100%}td,th{border-bottom:1px solid var(--line);padding:8px;text-align:left;
vertical-align:top}th{font-size:12px;text-transform:uppercase;letter-spacing:.04em;color:var(--muted)}
.seg{margin:0 0 12px;display:grid;grid-template-columns:70px 1fr;gap:8px}.who{font-weight:600}
.chip{display:inline-block;width:10px;height:10px;border-radius:50%;margin-right:6px}
.foot{margin-top:40px;color:var(--muted);font-size:12px;text-align:center}
@media print{body{margin:0}h2{break-after:avoid}.seg{break-inside:avoid}}
"""


def to_html(d: dict[str, Any], notes: bool = True, transcript: bool = True, timestamps: bool = True) -> str:
    e = html.escape
    out = [f"<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' "
           f"content='width=device-width,initial-scale=1'><title>{e(d['title'])}</title><style>{HTML_CSS}</style>"
           f"</head><body><h1>{e(d['title'])}</h1><div class='meta'>{e(_meta_line(d))}</div>"]
    n = d["notes"]
    if notes and n:
        for title, body in _notes_blocks(n):
            out.append(f"<h2>{e(title)}</h2>")
            if isinstance(body, str):
                out += [f"<p>{e(p)}</p>" for p in body.split("\n") if p.strip()]
            elif title == "Action items":
                out.append("<table><tr><th>Task</th><th>Owner</th><th>Due</th>" + ("<th>Time</th>" if timestamps else "") + "</tr>")
                for a in body:
                    out.append(f"<tr><td>{'✓ ' if a.get('done') else ''}{e(a['task'])}</td><td>{e(a.get('owner') or '')}</td>"
                               f"<td>{e(a.get('due') or '')}</td>" + (f"<td class='t'>{ts(a['t']) if a.get('t') is not None else ''}</td>" if timestamps else "") + "</tr>")
                out.append("</table>")
            elif title == "Chapters":
                items = []
                for c in body:
                    when = f"<span class='t'>{ts(c['t'])}</span> " if timestamps else ""
                    about = f" — {e(c['summary'])}" if c.get("summary") else ""
                    items.append(f"<li>{when}<b>{e(c['title'])}</b>{about}</li>")
                out.append("<ul>" + "".join(items) + "</ul>")
            else:
                out.append("<ul>" + "".join(f"<li>{e(x)}</li>" for x in body) + "</ul>")
    if transcript and d["segments"]:
        colors = {s["key"]: s["color"] for s in d["speakers"]}
        out.append("<h2>Transcript</h2>")
        for s in d["segments"]:
            col = colors.get(s["speaker"]) or "#7c3aed"
            out.append(f"<div class='seg'><div class='t'>{ts(s['start']) if timestamps else ''}</div><div>"
                       f"<span class='who' style='color:{e(col)}'>{e(s['name'])}:</span> {e(s['text'])}</div></div>")
    out.append(f"<div class='foot'>Created with QuickTranscriber {__version__} · processed privately on this computer</div>")
    out.append("</body></html>")
    return "".join(out)


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def export(meeting_id: str, fmt: str, notes: bool = True, transcript: bool = True, timestamps: bool = True,
           speakers: bool = True, include_audio: bool = True) -> tuple[bytes, str, str]:
    d = gather(meeting_id)
    if fmt not in FORMATS:
        raise ValueError(f"Unknown format {fmt}")
    kw = dict(notes=notes, transcript=transcript, timestamps=timestamps)
    if fmt == "txt":
        body = to_txt(d, **kw).encode("utf-8")
    elif fmt == "md":
        body = to_md(d, **kw).encode("utf-8")
    elif fmt == "docx":
        body = to_docx(d, **kw)
    elif fmt == "html":
        body = to_html(d, **kw).encode("utf-8")
    elif fmt == "srt":
        body = to_srt(d, speakers=speakers).encode("utf-8")
    elif fmt == "vtt":
        body = to_vtt(d, speakers=speakers).encode("utf-8")
    elif fmt == "json":
        body = to_json(d).encode("utf-8")
    else:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr(filename(d, "txt"), to_txt(d, **kw))
            z.writestr(filename(d, "md"), to_md(d, **kw))
            z.writestr(filename(d, "docx"), to_docx(d, **kw))
            z.writestr(filename(d, "html"), to_html(d, **kw))
            z.writestr(filename(d, "srt"), to_srt(d, speakers=speakers))
            z.writestr(filename(d, "vtt"), to_vtt(d, speakers=speakers))
            z.writestr(filename(d, "json"), to_json(d))
            if include_audio:
                mdir = paths.meeting_dir(meeting_id)
                originals = sorted(mdir.glob("original.*"))
                src = originals[0] if originals else mdir / "audio.wav"
                if src.exists():
                    z.write(src, filename(d, src.suffix.lstrip(".")), compress_type=zipfile.ZIP_STORED)
        body = buf.getvalue()
    return body, filename(d, fmt), FORMATS[fmt][1]
